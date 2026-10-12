"""
Fez ATS: Zoho Mail CV intake -> scoring -> HR pipeline dashboard -> stage emails.

Dashboard (HR):          streamlit run ats.py
Mailbox sync (cron):     python ats.py sync

Flow
1. Pull CV attachments (PDF/DOCX/DOC) from the Zoho mailbox over IMAP (read-only, mail is never altered).
2. Extract text, pick the role (title match, else Gemini/keywords) and score the CV 0-10 against that role's
   job description (entered once per role).
3. Score >= PASS_THRESHOLD (default 7) -> candidate is auto-moved to "Technical Assessment".
   Everything else stays in "Applied". Nobody is ever auto-rejected and no email is sent automatically.
4. HR reviews CVs inside the dashboard (no download needed), moves candidates between stages and can email
   the candidate at each move. All CVs stay available under "All CVs".
"""
from __future__ import annotations

import base64
import datetime as dt
import email
import functools
import hashlib
import hmac
import imaplib
import io
import json
import mimetypes
import os
import re
import smtplib
import sqlite3
import sys
import time
from contextlib import contextmanager
from email.header import decode_header, make_header
from email.message import EmailMessage
from email.utils import parseaddr, parsedate_to_datetime
from pathlib import Path

from dotenv import load_dotenv

load_dotenv()

# ----------------------------------------------------------------------------
# Config (all from .env)
# ----------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv("DATA_DIR", BASE_DIR / "data"))
CV_DIR = DATA_DIR / "cvs"
DB_PATH = DATA_DIR / "ats.db"
ASSETS_DIR = BASE_DIR / "assets"

ZOHO_EMAIL = os.getenv("ZOHO_EMAIL", "")
ZOHO_APP_PASSWORD = os.getenv("ZOHO_APP_PASSWORD", "")
IMAP_HOST = os.getenv("ZOHO_IMAP_HOST", "imap.zoho.com")  # imap.zoho.eu if your account is on the EU data centre
SMTP_HOST = os.getenv("ZOHO_SMTP_HOST", "smtp.zoho.com")
SMTP_PORT = int(os.getenv("ZOHO_SMTP_PORT", "465"))
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY", "")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
SYNC_DAYS = int(os.getenv("SYNC_DAYS", "7"))
PASS_THRESHOLD = float(os.getenv("PASS_THRESHOLD", "7"))
HR_SIGNATURE = os.getenv("HR_SIGNATURE", "HR Team, Fez Delivery Co.")
MAX_CV_BYTES = 15 * 1024 * 1024

# Logins: one HR user now; fill HR2_* in .env to add a second one later.
HR_USERS = {
    u: p
    for u, p in (
        (os.getenv("HR_USERNAME", "").strip(), os.getenv("HR_PASSWORD", "")),
        (os.getenv("HR2_USERNAME", "").strip(), os.getenv("HR2_PASSWORD", "")),
    )
    if u and p
}

# Branding: drop the logo into assets/ as logo.svg|png|jpg|webp and set the hex codes in .env.
BRAND_PRIMARY = os.getenv("BRAND_PRIMARY", "#45C734")   # Fez green
BRAND_ACCENT = os.getenv("BRAND_ACCENT", "#00243D")     # deep navy
BRAND_DARK = os.getenv("BRAND_DARK", "#00243D")         # deep navy (text and dark surfaces)
BRAND_BG = os.getenv("BRAND_BG", "#F7FAF6")             # light green tint
BRAND_WHITE = os.getenv("BRAND_WHITE", "#FFFFFF")
BRAND_MUTED = os.getenv("BRAND_MUTED", "#64748B")
BRAND_BORDER = os.getenv("BRAND_BORDER", "#E2E8E5")

STAGES = ["Applied", "Technical Assessment", "Interview", "People and Culture Assessment", "Offer Letter"]
REJECTED = "Rejected"
ALL_STAGES = STAGES + [REJECTED]
AUTO_STAGE = "Technical Assessment"
CV_EXTS = {".pdf", ".docx", ".doc"}

COMMON_SKILLS = [
    "excel", "sql", "python", "power bi", "tableau", "javascript", "java", "react", "node", "php",
    "communication", "customer service", "sales", "marketing", "accounting", "bookkeeping", "payroll",
    "logistics", "supply chain", "fleet management", "dispatch", "warehouse", "inventory", "procurement",
    "project management", "data analysis", "negotiation", "leadership", "social media", "seo",
    "recruitment", "hr", "audit", "tax", "ms office", "word", "powerpoint", "driving", "operations",
]

SIG = "{signature}"
DEFAULT_TEMPLATES = {
    "Applied": (
        "Your application to Fez Delivery Co.",
        "Dear {name},\n\nThank you for applying for the {role} position at Fez Delivery Co. "
        "We have received your application and will be in touch.\n\nKind regards,\n" + SIG,
    ),
    "Technical Assessment": (
        "Technical assessment: {role} at Fez Delivery Co.",
        "Dear {name},\n\nThank you for applying for the {role} position at Fez Delivery Co. "
        "We were impressed by your profile and would like to invite you to a technical assessment.\n\n"
        "Date: {date}\nTime: {time}\nVenue / Link: {venue}\n\n"
        "Please reply to confirm your availability.\n\nKind regards,\n" + SIG,
    ),
    "Interview": (
        "Interview invitation: {role} at Fez Delivery Co.",
        "Dear {name},\n\nCongratulations on progressing in our hiring process for the {role} position. "
        "We would like to invite you to an interview.\n\n"
        "Date: {date}\nTime: {time}\nVenue / Link: {venue}\n\n"
        "Please reply to confirm your availability and bring a valid ID.\n\nKind regards,\n" + SIG,
    ),
    "People and Culture Assessment": (
        "People and Culture assessment: {role} at Fez Delivery Co.",
        "Dear {name},\n\nYou have progressed to the People and Culture assessment for the {role} position.\n\n"
        "Date: {date}\nTime: {time}\nVenue / Link: {venue}\n\n"
        "Please reply to confirm your availability.\n\nKind regards,\n" + SIG,
    ),
    "Offer Letter": (
        "Offer of employment: {role} at Fez Delivery Co.",
        "Dear {name},\n\nWe are pleased to offer you the {role} position at Fez Delivery Co. "
        "Please find your offer letter attached. Kindly review it and reply to confirm your acceptance.\n\n"
        "Kind regards,\n" + SIG,
    ),
    "Rejected": (
        "Your application to Fez Delivery Co.",
        "Dear {name},\n\nThank you for your interest in the {role} position at Fez Delivery Co. and for the time "
        "you invested in applying. After careful consideration we will not be moving forward with your "
        "application at this time. We wish you every success.\n\nKind regards,\n" + SIG,
    ),
}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
EMAIL_FULL = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")
PHONE_RE = re.compile(r"(?:\+?234|0)[789][01]\d[\s\-]?\d{3}[\s\-]?\d{4}|\+?\d[\d\s\-()]{8,16}\d")
YEARS_RE = re.compile(r"(\d{1,2})\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:work\s+)?experience", re.I)


# ----------------------------------------------------------------------------
# Database
# ----------------------------------------------------------------------------
@contextmanager
def db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=30, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA journal_mode=WAL")
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def init_db() -> None:
    CV_DIR.mkdir(parents=True, exist_ok=True)
    with db() as c:
        c.executescript(
            """
            CREATE TABLE IF NOT EXISTS roles (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT UNIQUE NOT NULL,
                description TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS processed_mail (message_id TEXT PRIMARY KEY);
            CREATE TABLE IF NOT EXISTS candidates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                file_hash TEXT UNIQUE NOT NULL,
                received_at TEXT,
                sender_name TEXT,
                sender_email TEXT,
                subject TEXT,
                role TEXT,
                filename TEXT,
                file_path TEXT,
                name TEXT,
                email TEXT,
                phone TEXT,
                years_experience REAL,
                job_titles TEXT,
                skills TEXT,
                summary TEXT,
                missing_keywords TEXT,
                score REAL,
                parse_method TEXT,
                stage TEXT DEFAULT 'Applied',
                auto_advanced INTEGER DEFAULT 0,
                notes TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                stage_updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS stage_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                candidate_id INTEGER, from_stage TEXT, to_stage TEXT, by_user TEXT, note TEXT,
                at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS email_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                candidate_id INTEGER, stage TEXT, sent_to TEXT, subject TEXT, body TEXT, by_user TEXT,
                at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS templates (
                stage TEXT PRIMARY KEY, subject TEXT, body TEXT
            );
            """
        )
        for stage, (subj, body) in DEFAULT_TEMPLATES.items():
            c.execute("INSERT OR IGNORE INTO templates(stage, subject, body) VALUES (?,?,?)", (stage, subj, body))


def get_roles() -> list[sqlite3.Row]:
    with db() as c:
        return c.execute("SELECT * FROM roles ORDER BY title").fetchall()


def render_template(text: str, values: dict) -> str:
    for k, v in values.items():
        text = text.replace("{" + k + "}", str(v))
    return text


# ----------------------------------------------------------------------------
# Text extraction, role matching, scoring
# ----------------------------------------------------------------------------
def extract_text(data: bytes, ext: str) -> str:
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
        if ext == ".docx":
            import docx

            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            return "\n".join(parts).strip()
    except Exception:
        return ""
    return ""  # legacy .doc cannot be parsed; the file is still saved for HR


STOP = {"with", "that", "this", "have", "will", "from", "your", "their", "able", "work", "must", "about",
        "role", "team", "years", "experience", "and", "the", "for", "are", "you", "our", "who", "also"}


def keyword_score(text: str, jd: str) -> tuple[float | None, list[str]]:
    """Crude fallback: share of job-description keywords found in the CV, scaled 0-10."""
    if not (jd or "").strip():
        return None, []
    words = {w for w in re.findall(r"[a-zA-Z+#]{4,}", jd.lower()) if w not in STOP}
    if not words:
        return None, []
    low = text.lower()
    hit = {w for w in words if w in low}
    return round(10 * len(hit) / len(words), 1), sorted(words - hit)[:15]


def detect_role_by_title(subject: str, fname: str, body: str, roles) -> sqlite3.Row | None:
    hay = f"{subject} {fname} {body[:1500]}".lower()
    for r in sorted(roles, key=lambda r: -len(r["title"])):  # longest title first ("Senior X" before "X")
        if r["title"].lower() in hay:
            return r
    return None


def regex_parse(text: str, pool, fallback_name: str, fallback_email: str) -> dict:
    low = text.lower()
    emails = EMAIL_RE.findall(text)
    phones = PHONE_RE.findall(text)
    years = [int(y) for y in YEARS_RE.findall(text)]
    skills = [s for s in COMMON_SKILLS if re.search(r"\b" + re.escape(s) + r"\b", low)]
    first_line = next((ln.strip() for ln in text.splitlines() if 2 < len(ln.strip()) < 50), "")
    best_role, best_score, best_missing = "Unassigned", None, []
    for r in pool:
        sc, miss = keyword_score(text, r["description"] or "")
        if sc is not None and (best_score is None or sc > best_score):
            best_role, best_score, best_missing = r["title"], sc, miss
    if len(pool) > 1 and (best_score is None or best_score < 3):  # no convincing role among many
        best_role, best_score, best_missing = "Unassigned", None, []
    if not text.strip():  # unreadable (scanned) CV: never score it, HR reviews it by eye
        best_score, best_missing = None, []
    return {
        "name": first_line if first_line and not EMAIL_RE.search(first_line) else fallback_name,
        "email": emails[0] if emails else fallback_email,
        "phone": phones[0].strip() if phones else "",
        "years_experience": max(years) if years else None,
        "job_titles": [],
        "skills": skills,
        "summary": " ".join(text.split())[:300],
        "role": best_role,
        "match_score": best_score,
        "missing_keywords": best_missing,
    }


def gemini_parse(text: str, pool, subject: str = "", body: str = "") -> dict | None:
    if not GOOGLE_API_KEY or not text.strip():
        return None
    roles_block = "\n\n".join(f"ROLE: {r['title']}\nJOB DESCRIPTION: {(r['description'] or 'N/A')[:1500]}" for r in pool[:30])
    prompt = (
        "You are a strict ATS resume parser and screener. Return ONLY a JSON object with keys: "
        "name (string), email (string), phone (string), years_experience (number or null, total professional years), "
        "job_titles (list of strings, most recent first), skills (list of short strings), summary (max 40 words), "
        "role (the EXACT ROLE title from the list below that this candidate applied for or fits best, judging by the "
        "email subject, email body and CV; use \"Unassigned\" if none fits), "
        "match_score (number 0-10 with one decimal, scoring the CV against that role's job description: "
        "5 = partial fit, 7 = meets most requirements and is worth testing, 9+ = exceptional; null if role is Unassigned), "
        "missing_keywords (list of important job-description requirements absent from the CV).\n\n"
        f"ROLES:\n{roles_block or 'N/A'}\n\nEMAIL SUBJECT: {subject[:200]}\nEMAIL BODY: {body[:1000]}\n\nRESUME:\n{text[:15000]}"
    )
    for attempt in (1, 2):
        try:
            from google import genai

            client = genai.Client(api_key=GOOGLE_API_KEY)
            resp = client.models.generate_content(
                model=GEMINI_MODEL, contents=prompt, config={"response_mime_type": "application/json"}
            )
            data = json.loads(resp.text)
            if isinstance(data, list) and data:
                data = data[0]
            if isinstance(data, dict):
                titles = {r["title"].lower(): r["title"] for r in pool}
                data["role"] = titles.get(str(data.get("role", "")).strip().lower(), "Unassigned")
                if data["role"] == "Unassigned":
                    data["match_score"] = None
                return data
        except Exception:
            if attempt == 1:
                time.sleep(2)
    return None


def to_float(v) -> float | None:
    try:
        return float(v) if v is not None else None
    except (TypeError, ValueError):
        return None


def clamp_score(v) -> float | None:
    f = to_float(v)
    return None if f is None else round(max(0.0, min(10.0, f)), 1)


def join_list(v) -> str:
    if isinstance(v, list):
        return ", ".join(str(x).strip() for x in v if str(x).strip())
    return str(v or "")


def analyse(text: str, pool, subject: str, body: str, s_name: str, s_email: str, fname: str) -> tuple[dict, str]:
    parsed = gemini_parse(text, pool, subject, body)
    if parsed:
        return parsed, "gemini"
    parsed = regex_parse(text, pool, s_name or Path(fname).stem, s_email)
    return parsed, ("keywords" if text else "no-text (review manually)")


# ----------------------------------------------------------------------------
# Stage movement and email
# ----------------------------------------------------------------------------
def move_stage(cid: int, new_stage: str, by: str, note: str = "", auto: bool = False) -> None:
    if new_stage not in ALL_STAGES:
        raise ValueError(f"Unknown stage: {new_stage}")
    with db() as c:
        row = c.execute("SELECT stage FROM candidates WHERE id=?", (cid,)).fetchone()
        if row is None:
            raise ValueError("Candidate not found")
        c.execute(
            "UPDATE candidates SET stage=?, stage_updated_at=CURRENT_TIMESTAMP, auto_advanced=? WHERE id=?",
            (new_stage, 1 if auto else 0, cid),
        )
        c.execute(
            "INSERT INTO stage_log(candidate_id, from_stage, to_stage, by_user, note) VALUES (?,?,?,?,?)",
            (cid, row["stage"], new_stage, by, note),
        )


def send_email(cid: int, stage: str, to: str, subject: str, body: str, by: str, attachment=None) -> None:
    """attachment: optional (filename, bytes). Raises on any problem; logs only after a successful send."""
    if not ZOHO_EMAIL or not ZOHO_APP_PASSWORD:
        raise RuntimeError("Set ZOHO_EMAIL and ZOHO_APP_PASSWORD in .env to send mail")
    to = (to or "").strip()
    if not EMAIL_FULL.match(to):
        raise ValueError(f"'{to}' is not a valid email address; fix it in the candidate details")
    msg = EmailMessage()
    msg["From"] = ZOHO_EMAIL
    msg["To"] = to
    msg["Subject"] = subject.replace("\n", " ").strip()
    msg.set_content(body)
    if attachment:
        fname, data = attachment
        mime = mimetypes.guess_type(fname)[0] or "application/octet-stream"
        msg.add_attachment(data, maintype=mime.split("/")[0], subtype=mime.split("/")[1], filename=fname)
    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as s:
        s.login(ZOHO_EMAIL, ZOHO_APP_PASSWORD)
        s.send_message(msg)
    with db() as c:
        c.execute(
            "INSERT INTO email_log(candidate_id, stage, sent_to, subject, body, by_user) VALUES (?,?,?,?,?,?)",
            (cid, stage, to, subject, body, by),
        )


# ----------------------------------------------------------------------------
# Ingest + mailbox sync
# ----------------------------------------------------------------------------
def dec(value: str | None) -> str:
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return value


def safe_name(s: str) -> str:
    return re.sub(r"[^\w.\- ]+", "_", s).strip() or "file"


def get_body(msg) -> str:
    for part in msg.walk():
        if part.get_content_type() == "text/plain" and not part.get_filename():
            try:
                payload = part.get_payload(decode=True) or b""
                return payload.decode(part.get_content_charset() or "utf-8", errors="ignore")[:3000]
            except Exception:
                continue
    return ""


def ingest_cv(c, payload: bytes, fname: str, subject: str, body: str, s_name: str, s_email: str,
              received: str, roles, log=print) -> str:
    """Parse, score, store one CV. Returns 'new' or 'duplicate'."""
    ext = Path(fname).suffix.lower()
    fhash = hashlib.sha256(payload).hexdigest()
    if c.execute("SELECT 1 FROM candidates WHERE file_hash=?", (fhash,)).fetchone():
        return "duplicate"

    text = extract_text(payload, ext)
    title_role = detect_role_by_title(subject, fname, body, roles)
    pool = [title_role] if title_role else list(roles)
    parsed, method = analyse(text, pool, subject, body, s_name, s_email, fname)

    role = parsed.get("role") or (title_role["title"] if title_role else "Unassigned")
    score = clamp_score(parsed.get("match_score"))
    folder = CV_DIR / safe_name(role) / received[:10]
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{fhash[:8]}_{safe_name(fname)}"
    path.write_bytes(payload)

    cur = c.execute(
        """INSERT INTO candidates
        (file_hash, received_at, sender_name, sender_email, subject, role, filename, file_path,
         name, email, phone, years_experience, job_titles, skills, summary, missing_keywords, score, parse_method)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            fhash, received, s_name, s_email, subject, role, fname, str(path),
            str(parsed.get("name") or s_name or Path(fname).stem)[:120],
            str(parsed.get("email") or s_email)[:120],
            str(parsed.get("phone") or "")[:40],
            to_float(parsed.get("years_experience")),
            join_list(parsed.get("job_titles")),
            join_list(parsed.get("skills")),
            str(parsed.get("summary") or "")[:500],
            join_list(parsed.get("missing_keywords")),
            score, method,
        ),
    )
    cid = cur.lastrowid
    if score is not None and score >= PASS_THRESHOLD:
        c.execute("UPDATE candidates SET stage=?, auto_advanced=1 WHERE id=?", (AUTO_STAGE, cid))
        c.execute(
            "INSERT INTO stage_log(candidate_id, from_stage, to_stage, by_user, note) VALUES (?,?,?,?,?)",
            (cid, "Applied", AUTO_STAGE, "system", f"Auto-advanced: score {score}/10 >= {PASS_THRESHOLD:g}"),
        )
    log(f"+ {fname} | {role} | score {score} | {method}")
    return "new"


def sync_mailbox(days: int = SYNC_DAYS, log=print) -> dict:
    if not ZOHO_EMAIL or not ZOHO_APP_PASSWORD:
        raise RuntimeError("Set ZOHO_EMAIL and ZOHO_APP_PASSWORD in .env")
    init_db()
    roles = get_roles()
    stats = {"emails": 0, "new_cvs": 0, "duplicates": 0, "skipped": 0, "errors": 0}

    conn = imaplib.IMAP4_SSL(IMAP_HOST, 993)
    try:
        conn.login(ZOHO_EMAIL, ZOHO_APP_PASSWORD)
        conn.select("INBOX", readonly=True)  # read-only: never marks mail as read
        since = (dt.date.today() - dt.timedelta(days=days)).strftime("%d-%b-%Y")
        typ, data = conn.search(None, "SINCE", since)
        if typ != "OK":
            raise RuntimeError("IMAP search failed")
        nums = data[0].split()
        log(f"{len(nums)} emails in the last {days} days")

        with db() as c:
            for num in nums:
                typ, hdr = conn.fetch(num, "(BODY.PEEK[HEADER.FIELDS (MESSAGE-ID)])")
                mid = ""
                if typ == "OK" and hdr and isinstance(hdr[0], tuple):
                    mid = hdr[0][1].decode(errors="ignore").split(":", 1)[-1].strip()
                mid = mid or f"no-id-{num.decode()}"
                if c.execute("SELECT 1 FROM processed_mail WHERE message_id=?", (mid,)).fetchone():
                    stats["skipped"] += 1
                    continue

                typ, msgdata = conn.fetch(num, "(BODY.PEEK[])")
                if typ != "OK" or not msgdata or not isinstance(msgdata[0], tuple):
                    continue
                msg = email.message_from_bytes(msgdata[0][1])
                stats["emails"] += 1

                subject = dec(msg.get("Subject"))
                body = get_body(msg)
                s_name, s_email = parseaddr(dec(msg.get("From")))
                try:
                    received = parsedate_to_datetime(msg.get("Date")).strftime("%Y-%m-%d %H:%M")
                except Exception:
                    received = dt.datetime.now().strftime("%Y-%m-%d %H:%M")

                failed = False
                for part in msg.walk():
                    fname = dec(part.get_filename())
                    if not fname or Path(fname).suffix.lower() not in CV_EXTS:
                        continue
                    payload = part.get_payload(decode=True)
                    if not payload or len(payload) > MAX_CV_BYTES:
                        continue
                    try:
                        res = ingest_cv(c, payload, fname, subject, body, s_name, s_email, received, roles, log)
                        stats["new_cvs" if res == "new" else "duplicates"] += 1
                    except Exception as e:  # one bad CV must not stop the run
                        failed = True
                        stats["errors"] += 1
                        log(f"! {fname}: {e}")
                if not failed:  # failed emails are retried on the next run
                    c.execute("INSERT OR IGNORE INTO processed_mail(message_id) VALUES (?)", (mid,))
                c.commit()
    finally:
        try:
            conn.logout()
        except Exception:
            pass
    return stats


def rescore(role_title: str | None = None, log=print) -> dict:
    """Re-analyse stored CVs against the current job descriptions. Never moves anyone backwards."""
    init_db()
    roles = get_roles()
    by_title = {r["title"]: r for r in roles}
    stats = {"rescored": 0, "advanced": 0, "missing_file": 0}
    with db() as c:
        sql = "SELECT * FROM candidates" + (" WHERE role=?" if role_title else "")
        rows = c.execute(sql, (role_title,) if role_title else ()).fetchall()
        for r in rows:
            p = Path(r["file_path"] or "")
            if not r["file_path"] or not p.is_file():
                stats["missing_file"] += 1
                continue
            text = extract_text(p.read_bytes(), p.suffix.lower())
            pool = [by_title[r["role"]]] if r["role"] in by_title else list(roles)
            parsed, method = analyse(text, pool, r["subject"] or "", "", r["sender_name"] or "", r["sender_email"] or "", r["filename"] or p.name)
            score = clamp_score(parsed.get("match_score"))
            new_role = parsed.get("role") or r["role"]
            c.execute(
                "UPDATE candidates SET score=?, role=?, missing_keywords=?, parse_method=? WHERE id=?",
                (score, new_role if r["role"] == "Unassigned" else r["role"],
                 join_list(parsed.get("missing_keywords")), method, r["id"]),
            )
            stats["rescored"] += 1
            if r["stage"] == "Applied" and score is not None and score >= PASS_THRESHOLD:
                c.execute("UPDATE candidates SET stage=?, auto_advanced=1, stage_updated_at=CURRENT_TIMESTAMP WHERE id=?", (AUTO_STAGE, r["id"]))
                c.execute(
                    "INSERT INTO stage_log(candidate_id, from_stage, to_stage, by_user, note) VALUES (?,?,?,?,?)",
                    (r["id"], "Applied", AUTO_STAGE, "system", f"Auto-advanced after re-score: {score}/10"),
                )
                stats["advanced"] += 1
    log(str(stats))
    return stats


# ----------------------------------------------------------------------------
# Branding helpers
# ----------------------------------------------------------------------------
def logo_data_uri() -> str | None:
    for ext in ("svg", "png", "jpg", "jpeg", "webp"):
        p = ASSETS_DIR / f"logo.{ext}"
        if p.exists():
            mime = "image/svg+xml" if ext == "svg" else f"image/{'jpeg' if ext in ('jpg', 'jpeg') else ext}"
            return f"data:{mime};base64,{base64.b64encode(p.read_bytes()).decode()}"
    return None


def logo_html(height: int = 44) -> str:
    uri = logo_data_uri()
    if uri:
        return f'<img src="{uri}" style="height:{height}px;width:auto;display:block;" alt="Fez Delivery Co.">'
    return f'<span style="font-weight:800;font-size:{int(height*0.6)}px;color:{BRAND_PRIMARY};">Fez Delivery Co.</span>'


def ensure_light_theme() -> None:
    """Streamlit follows the viewer's OS dark mode. The Fez design is light, so pin the theme to light.
    Written once into .streamlit/config.toml (next to where the app is started); applies after a restart."""
    try:
        cfg = Path.cwd() / ".streamlit" / "config.toml"
        cfg.parent.mkdir(exist_ok=True)
        text = cfg.read_text() if cfg.exists() else ""
        if "[theme]" in text:
            return
        block = (
            '\n[theme]\nbase = "light"\n'
            f'primaryColor = "{BRAND_PRIMARY}"\nbackgroundColor = "{BRAND_BG}"\n'
            f'secondaryBackgroundColor = "#EDF3EA"\ntextColor = "{BRAND_DARK}"\n'
        )
        cfg.write_text(text + block)
    except Exception:
        pass  # read-only folder: the CSS below still keeps the page readable


def brand_css() -> str:
    d, w, b = BRAND_DARK, BRAND_WHITE, BRAND_BORDER
    return f"""<style>
    :root {{ color-scheme: light; }}
    .stApp {{ background:{BRAND_BG}; color:{d}; }}
    [data-testid="stHeader"] {{ background:{BRAND_BG}; }}
    [data-testid="stSidebar"] {{ background:#EDF3EA; }}

    /* text stays dark whatever the viewer's OS theme is */
    .stApp h1, .stApp h2, .stApp h3, .stApp h4 {{ color:{d}; }}
    .stApp p, .stApp li, .stApp label, .stApp span, .stApp small, .stApp summary,
    .stApp [data-testid="stMarkdownContainer"], .stApp [data-testid="stWidgetLabel"],
    .stApp [data-testid="stCaptionContainer"], .stApp [data-testid="stMetricValue"],
    .stApp [data-testid="stMetricLabel"], .stApp [data-testid="stExpander"] summary {{ color:{d}; }}
    .stApp [data-testid="stCaptionContainer"] {{ color:{BRAND_MUTED}; }}

    /* inputs, dropdowns, uploaders */
    .stApp input, .stApp textarea, .stApp [data-baseweb="select"] > div,
    .stApp [data-baseweb="input"], .stApp [data-baseweb="textarea"], .stApp [data-testid="stFileUploaderDropzone"] {{
        background:{w} !important; color:{d} !important; border-color:{b}; }}
    .stApp [data-baseweb="select"] * {{ color:{d}; }}
    [data-baseweb="popover"] ul, [data-baseweb="popover"] li, [data-baseweb="menu"], [role="listbox"],
    [data-baseweb="calendar"], [data-baseweb="calendar"] * {{ background:{w}; color:{d}; }}

    /* buttons */
    .stApp .stButton > button, .stApp .stDownloadButton > button, .stApp [data-testid="stBaseButton-secondary"] {{
        background:{w}; color:{d}; border:1px solid {b}; }}
    .stApp .stButton > button *, .stApp .stDownloadButton > button * {{ color:inherit; }}
    .stApp .stButton > button:hover, .stApp .stDownloadButton > button:hover {{
        border-color:{BRAND_PRIMARY}; color:{d}; }}
    .stApp .stButton > button[kind="primary"], .stApp .stDownloadButton > button[kind="primary"],
    .stApp div[data-testid="stFormSubmitButton"] > button {{
        background:{BRAND_PRIMARY}; border:1px solid {BRAND_PRIMARY}; color:{d}; font-weight:700; }}
    .stApp .stButton > button[kind="primary"]:hover, .stApp div[data-testid="stFormSubmitButton"] > button:hover {{
        background:{d}; border-color:{d}; color:#fff; }}
    .stApp .stButton > button[kind="primary"]:hover * {{ color:#fff; }}

    /* header, metric cards, tabs */
    .fez-header {{ display:flex; align-items:center; gap:16px; padding:10px 18px; margin-bottom:14px;
        background:{w}; border:1px solid {b}; border-radius:10px; border-bottom:4px solid {BRAND_PRIMARY}; }}
    .fez-header .t {{ font-size:20px; font-weight:700; color:{d}; }}
    .stApp div[data-testid="stMetric"] {{ background:{w}; border:1px solid {b}; border-radius:10px;
        padding:10px 14px; border-left:5px solid {BRAND_PRIMARY}; }}
    .stApp button[data-baseweb="tab"] {{ color:{BRAND_MUTED}; }}
    .stApp button[data-baseweb="tab"][aria-selected="true"], .stApp button[data-baseweb="tab"][aria-selected="true"] * {{
        color:{d}; font-weight:700; }}
    .stApp div[data-baseweb="tab-highlight"] {{ background-color:{BRAND_PRIMARY}; height:3px; }}
    .stApp [data-testid="stExpander"] {{ background:{w}; border:1px solid {b}; border-radius:8px; }}

    /* chips: green is bright, so its text is navy */
    .stApp .chip {{ display:inline-block; padding:3px 11px; border-radius:999px; font-size:13px; font-weight:700;
        background:{BRAND_PRIMARY}; color:{d} !important; margin-right:6px; }}
    .stApp .chip.rej {{ background:#B42318; color:#fff !important; }}
    .stApp .chip.good {{ background:{BRAND_PRIMARY}; color:{d} !important; }}
    .stApp .chip.mid {{ background:#F5B83D; color:{d} !important; }}
    .stApp .chip.low {{ background:#E2E8F0; color:#334155 !important; }}
    </style>"""


def score_chip(score) -> str:
    if score is None or score != score:
        return '<span class="chip low">Not scored</span>'
    cls = "good" if score >= PASS_THRESHOLD else ("mid" if score >= PASS_THRESHOLD - 2 else "low")
    return f'<span class="chip {cls}">Score {score:g}/10</span>'


@functools.lru_cache(maxsize=24)
def render_pdf_pages(path: str, mtime: float, max_pages: int = 8) -> tuple[list[bytes], int]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(path)
    try:
        total = len(pdf)
        out = []
        for i in range(min(total, max_pages)):
            page = pdf[i]
            img = page.render(scale=1.6).to_pil()
            page.close()
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            out.append(buf.getvalue())
        return out, total
    finally:
        pdf.close()


# ----------------------------------------------------------------------------
# Dashboard
# ----------------------------------------------------------------------------
def run_dashboard() -> None:
    import pandas as pd
    import streamlit as st

    st.set_page_config(page_title="Fez ATS", layout="wide")
    init_db()
    ensure_light_theme()
    st.markdown(brand_css(), unsafe_allow_html=True)

    # ---------------- login ----------------
    if not HR_USERS:
        st.error("No login configured. Set HR_USERNAME and HR_PASSWORD in .env and restart.")
        st.stop()
    if "user" not in st.session_state:
        _, mid, _ = st.columns([1, 1.2, 1])
        with mid:
            st.markdown(f'<div class="fez-header">{logo_html(48)}<span class="t">Recruitment ATS</span></div>', unsafe_allow_html=True)
            with st.form("login"):
                u = st.text_input("Username")
                p = st.text_input("Password", type="password")
                if st.form_submit_button("Sign in", type="primary"):
                    ok = any(hmac.compare_digest(u.strip(), name) and hmac.compare_digest(p, pw) for name, pw in HR_USERS.items())
                    if ok:
                        st.session_state["user"] = u.strip()
                        st.rerun()
                    st.session_state["fails"] = st.session_state.get("fails", 0) + 1
                    time.sleep(min(2 ** st.session_state["fails"], 10))
                    st.error("Wrong username or password")
        st.stop()
    user = st.session_state["user"]

    st.markdown(f'<div class="fez-header">{logo_html(48)}<span class="t">Recruitment ATS</span></div>', unsafe_allow_html=True)
    if st.session_state.get("flash"):
        st.success(st.session_state.pop("flash"))

    # ---------------- sidebar ----------------
    with st.sidebar:
        st.markdown(f"Signed in as **{user}**")
        if st.button("Sign out"):
            st.session_state.clear()
            st.rerun()
        st.divider()
        st.subheader("Mailbox sync")
        days = st.number_input("Look back (days)", 1, 90, SYNC_DAYS)
        if st.button("Sync Zoho inbox now", type="primary"):
            try:
                with st.spinner("Reading mailbox and scoring CVs..."):
                    stats = sync_mailbox(int(days), log=lambda m: None)
                st.session_state["flash"] = (
                    f"{stats['new_cvs']} new CVs, {stats['duplicates']} duplicates, "
                    f"{stats['errors']} errors, {stats['skipped']} emails already processed"
                )
                st.rerun()
            except Exception as e:
                st.error(f"Sync failed: {e}")
        st.caption(f"Scoring: {'Gemini (keyword fallback)' if GOOGLE_API_KEY else 'keywords only (no GOOGLE_API_KEY)'}")
        st.caption(f"Auto-advance at score {PASS_THRESHOLD:g}/10 or higher")

    # ---------------- shared pieces ----------------
    def load_df() -> "pd.DataFrame":
        with db() as c:
            return pd.read_sql_query("SELECT * FROM candidates ORDER BY (score IS NULL), score DESC, received_at DESC", c)

    def cv_viewer(row, ns: str) -> None:
        p = Path(row["file_path"] or "")
        if not row["file_path"] or not p.is_file():
            st.warning("CV file not found on the server.")
            return
        ext = p.suffix.lower()
        if ext == ".pdf":
            try:
                pages, total = render_pdf_pages(str(p), p.stat().st_mtime)
                for img in pages:
                    st.image(img)
                if total > len(pages):
                    st.caption(f"Showing the first {len(pages)} of {total} pages.")
            except Exception:
                st.warning("Could not render this PDF. Use the download button.")
        elif ext == ".docx":
            text = extract_text(p.read_bytes(), ext)
            with st.container(height=600):
                st.text(text or "(No readable text in this document.)")
        else:
            st.info("Old .doc files cannot be previewed. Use the download button.")
        st.download_button("Download original file", p.read_bytes(), p.name, key=f"{ns}_dl_{row['id']}")

    def candidate_panel(cid: int, ns: str) -> None:
        with db() as c:
            row = c.execute("SELECT * FROM candidates WHERE id=?", (cid,)).fetchone()
            logs = c.execute("SELECT * FROM stage_log WHERE candidate_id=? ORDER BY id DESC", (cid,)).fetchall()
            mails = c.execute("SELECT * FROM email_log WHERE candidate_id=? ORDER BY id DESC", (cid,)).fetchall()
        if row is None:
            st.warning("Candidate not found.")
            return
        chip_cls = "chip rej" if row["stage"] == REJECTED else "chip"
        st.markdown(
            f"### {row['name']}\n<span class='{chip_cls}'>{row['stage']}</span>{score_chip(row['score'])}"
            f"<span class='chip low'>{row['role']}</span>",
            unsafe_allow_html=True,
        )
        left, right = st.columns([3, 2], gap="large")

        with left:
            st.write(f"**Email:** {row['email'] or 'n/a'}  |  **Phone:** {row['phone'] or 'n/a'}")
            yrs = row["years_experience"]
            st.write(f"**Experience:** {'n/a' if yrs is None or yrs != yrs else f'{yrs:g} yrs'}  |  **Received:** {row['received_at']}")
            if row["job_titles"]:
                st.write(f"**Titles:** {row['job_titles']}")
            if row["skills"]:
                st.write(f"**Skills:** {row['skills']}")
            if row["summary"]:
                st.write(f"**Summary:** {row['summary']}")
            if row["missing_keywords"]:
                st.write(f"**Missing vs job description:** {row['missing_keywords']}")
            method = row["parse_method"] or ""
            if method != "gemini":
                st.caption(f"Scored with: {method}. Treat this score with extra caution.")
            if row["auto_advanced"]:
                st.caption("Moved to Technical Assessment automatically by score.")
            st.markdown("**CV**")
            cv_viewer(row, ns)

        with right:
            st.markdown("**Move candidate**")
            cur = row["stage"]
            idx = STAGES.index(cur) if cur in STAGES else 0
            default = STAGES[min(idx + 1, len(STAGES) - 1)] if cur in STAGES else "Applied"
            target = st.selectbox("Move to", ALL_STAGES, index=ALL_STAGES.index(default), key=f"{ns}_t_{cid}")
            can_email = bool(ZOHO_EMAIL and ZOHO_APP_PASSWORD)
            do_mail = st.checkbox("Email the candidate", value=False, key=f"{ns}_m_{cid}", disabled=not can_email,
                                  help=None if can_email else "Set ZOHO_EMAIL and ZOHO_APP_PASSWORD in .env")
            subject = body = ""
            attach = None
            if do_mail:
                with db() as c:
                    tpl = c.execute("SELECT subject, body FROM templates WHERE stage=?", (target,)).fetchone()
                needs_slot = target in ("Technical Assessment", "Interview", "People and Culture Assessment")
                d_val = t_val = v_val = ""
                if needs_slot:
                    d1, d2 = st.columns(2)
                    d_val = d1.date_input("Date", dt.date.today() + dt.timedelta(days=3), key=f"{ns}_d_{cid}").strftime("%A, %d %B %Y")
                    t_val = d2.text_input("Time", "10:00 AM", key=f"{ns}_ti_{cid}")
                    v_val = st.text_input("Venue / meeting link", "Fez HQ, Lagos", key=f"{ns}_v_{cid}")
                vals = {"name": row["name"], "role": row["role"], "date": d_val, "time": t_val, "venue": v_val, "signature": HR_SIGNATURE}
                sig = f"{cid}_{target}_{d_val}_{t_val}_{v_val}"  # editing the date/time refreshes the draft
                subject = st.text_input("Subject", render_template(tpl["subject"], vals), key=f"{ns}_s_{sig}")
                body = st.text_area("Message", render_template(tpl["body"], vals), height=260, key=f"{ns}_b_{sig}")
                st.caption(f"Sends to: {row['email'] or 'NO EMAIL ON FILE'}. Edit the date, time and venue first; changing them redraws the message.")
                up = st.file_uploader("Attachment (optional)" + (", e.g. the signed offer letter PDF" if target == "Offer Letter" else ""),
                                      key=f"{ns}_f_{cid}")
                if up is not None:
                    attach = (up.name, up.getvalue())
            note = st.text_input("Note for the history (optional)", key=f"{ns}_n_{cid}")
            if st.button("Confirm", type="primary", key=f"{ns}_go_{cid}", disabled=(target == cur and not do_mail)):
                try:
                    msg = []
                    if do_mail:
                        send_email(cid, target, row["email"], subject, body, user, attach)
                        msg.append(f"Email sent to {row['email']}")
                    if target != cur:
                        move_stage(cid, target, user, note)
                        msg.append(f"moved to {target}")
                    st.session_state["flash"] = f"{row['name']}: " + ", ".join(msg)
                    st.rerun()
                except Exception as e:
                    st.error(f"Not done: {e}")

            st.markdown("**HR notes**")
            notes = st.text_area("Notes", row["notes"] or "", key=f"{ns}_notes_{cid}", label_visibility="collapsed")
            if st.button("Save notes", key=f"{ns}_sn_{cid}"):
                with db() as c:
                    c.execute("UPDATE candidates SET notes=? WHERE id=?", (notes, cid))
                st.session_state["flash"] = "Notes saved"
                st.rerun()
            with st.expander(f"Stage history ({len(logs)})"):
                for lg in logs:
                    st.write(f"{lg['at']}: {lg['from_stage']} to **{lg['to_stage']}** by {lg['by_user']}" + (f" ({lg['note']})" if lg["note"] else ""))
            with st.expander(f"Emails sent ({len(mails)})"):
                for m in mails:
                    st.write(f"{m['at']}: **{m['subject']}** to {m['sent_to']} by {m['by_user']}")

    def pick_row(df: "pd.DataFrame", cols: list[str], key: str) -> int | None:
        ev = st.dataframe(df[cols], hide_index=True, on_select="rerun", selection_mode="single-row", key=f"tbl_{key}")
        rows = ev.selection.rows if ev is not None else []
        if rows:
            st.session_state[f"open_{key}"] = int(df.iloc[rows[0]]["id"])
        cid = st.session_state.get(f"open_{key}")
        return cid if cid in set(df["id"]) else None

    show_cols = ["id", "name", "role", "score", "stage", "received_at", "email", "phone", "years_experience"]

    tab_p, tab_a, tab_r, tab_t, tab_l = st.tabs(["Overview", "All CVs", "Roles and scoring", "Email templates", "Activity"])

    # ---------------- Pipeline ----------------
    with tab_p:
        df = load_df()
        counts = df["stage"].value_counts().to_dict() if not df.empty else {}
        cols = st.columns(len(ALL_STAGES))
        for col, s in zip(cols, ALL_STAGES):
            col.metric(s, counts.get(s, 0))
        stage_sel = st.radio("Stage", ALL_STAGES, horizontal=True, key="pipe_stage",
                             format_func=lambda s: f"{s} ({counts.get(s, 0)})", label_visibility="collapsed")
        v = df[df["stage"] == stage_sel] if not df.empty else df
        roles_here = sorted(v["role"].dropna().unique()) if not v.empty else []
        rf = st.multiselect("Filter by role", roles_here, key="pipe_role")
        if rf:
            v = v[v["role"].isin(rf)]
        if v.empty:
            st.info("No candidates in this stage.")
        else:
            st.caption("Click a row to open the candidate. Sorted by score, highest first.")
            cid = pick_row(v.reset_index(drop=True), show_cols, "pipe")
            if cid:
                st.divider()
                candidate_panel(cid, "p")

    # ---------------- All CVs ----------------
    with tab_a:
        df = load_df()
        if df.empty:
            st.info("No CVs yet. Add roles, then click 'Sync Zoho inbox now'.")
        else:
            f1, f2, f3, f4 = st.columns(4)
            role_f = f1.multiselect("Role", sorted(df["role"].dropna().unique()), key="all_role")
            stage_f = f2.multiselect("Stage", ALL_STAGES, key="all_stage")
            min_score = f3.slider("Min score", 0.0, 10.0, 0.0, 0.5, key="all_score")
            min_years = f4.number_input("Min years experience", 0.0, 50.0, 0.0, 0.5, key="all_years")
            q = st.text_input("Search name, skills, titles, summary, filename (comma separated, all must match)", key="all_q")
            v = df.copy()
            if role_f:
                v = v[v["role"].isin(role_f)]
            if stage_f:
                v = v[v["stage"].isin(stage_f)]
            if min_score:
                v = v[v["score"].fillna(-1) >= min_score]
            if min_years:
                v = v[v["years_experience"].fillna(-1) >= min_years]
            blob = (v["name"].fillna("") + " " + v["skills"].fillna("") + " " + v["job_titles"].fillna("") + " "
                    + v["summary"].fillna("") + " " + v["filename"].fillna("") + " " + v["email"].fillna("")).str.lower()
            for k in [x.strip().lower() for x in q.split(",") if x.strip()]:
                v = v[blob.loc[v.index].str.contains(re.escape(k))]
            st.caption(f"{len(v)} of {len(df)} CVs (every submission is kept, whatever its score or stage)")
            v = v.reset_index(drop=True)
            if not v.empty:
                st.download_button("Download this list (CSV)", v[show_cols].to_csv(index=False).encode(), "candidates.csv", "text/csv")
                cid = pick_row(v, show_cols, "all")
                if cid:
                    st.divider()
                    candidate_panel(cid, "a")

    # ---------------- Roles and scoring ----------------
    with tab_r:
        st.subheader("Add or update a role")
        st.caption("Enter the job description once per role. Every incoming CV for that role is scored against it automatically.")
        with st.form("role_form", clear_on_submit=True):
            title = st.text_input("Role title (e.g. Dispatch Rider, Data Analyst)")
            desc = st.text_area("Job description: responsibilities, required skills, experience, qualifications", height=200)
            if st.form_submit_button("Save role", type="primary") and title.strip():
                with db() as c:
                    c.execute(
                        "INSERT INTO roles(title, description) VALUES (?,?) "
                        "ON CONFLICT(title) DO UPDATE SET description=excluded.description",
                        (title.strip(), desc.strip()),
                    )
                st.success("Saved. New CVs use it from now on. Use the re-score box below for CVs already received.")
        roles = get_roles()
        if roles:
            st.dataframe(pd.DataFrame([dict(r) for r in roles])[["title", "description"]], hide_index=True)
            st.subheader("Re-score existing CVs")
            st.caption("Use after editing a job description. Updates scores and moves 'Applied' candidates who now reach the threshold. Nobody is moved backwards.")
            rs = st.selectbox("Role", ["All roles"] + [r["title"] for r in roles], key="rs_role")
            if st.button("Re-score now"):
                with st.spinner("Re-scoring..."):
                    res = rescore(None if rs == "All roles" else rs, log=lambda m: None)
                st.session_state["flash"] = f"Re-scored {res['rescored']} CVs, {res['advanced']} moved to {AUTO_STAGE}"
                st.rerun()
        else:
            st.info("No roles yet. CVs are filed under 'Unassigned' and are not scored until you add roles.")

    # ---------------- Email templates ----------------
    with tab_t:
        st.caption("Placeholders: {name} {role} {date} {time} {venue} {signature}. These fill the draft HR sees before sending.")
        ts = st.selectbox("Template for stage", ALL_STAGES, key="tpl_stage")
        with db() as c:
            tpl = c.execute("SELECT subject, body FROM templates WHERE stage=?", (ts,)).fetchone()
        n_sub = st.text_input("Subject", tpl["subject"], key=f"tpl_s_{ts}")
        n_body = st.text_area("Body", tpl["body"], height=300, key=f"tpl_b_{ts}")
        c1, c2 = st.columns(2)
        if c1.button("Save template", type="primary"):
            with db() as c:
                c.execute("UPDATE templates SET subject=?, body=? WHERE stage=?", (n_sub, n_body, ts))
            st.session_state["flash"] = "Template saved"
            st.rerun()
        if c2.button("Reset to default"):
            d = DEFAULT_TEMPLATES[ts]
            with db() as c:
                c.execute("UPDATE templates SET subject=?, body=? WHERE stage=?", d + (ts,))
            for k in (f"tpl_s_{ts}", f"tpl_b_{ts}"):
                st.session_state.pop(k, None)
            st.rerun()

    # ---------------- Activity ----------------
    with tab_l:
        with db() as c:
            sl = pd.read_sql_query(
                "SELECT l.at, c.name, c.role, l.from_stage, l.to_stage, l.by_user, l.note FROM stage_log l "
                "LEFT JOIN candidates c ON c.id=l.candidate_id ORDER BY l.id DESC LIMIT 200", c)
            el = pd.read_sql_query(
                "SELECT e.at, c.name, e.stage, e.sent_to, e.subject, e.by_user FROM email_log e "
                "LEFT JOIN candidates c ON c.id=e.candidate_id ORDER BY e.id DESC LIMIT 200", c)
        st.subheader("Stage changes")
        st.dataframe(sl, hide_index=True)
        st.subheader("Emails sent")
        st.dataframe(el, hide_index=True)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "sync":
        init_db()
        print(sync_mailbox(log=print))
    elif len(sys.argv) > 1 and sys.argv[1] == "rescore":
        print(rescore())
    else:
        run_dashboard()